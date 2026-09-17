from __future__ import annotations

from types import SimpleNamespace

import openpyxl
import pytest

from src.assembleur_chemins_export import (
    CheminsExportError,
    PointsExportOptions,
    TripletsExportOptions,
    build_points_rows,
    build_triplets_rows,
    export_chemins_xlsx,
    unique_preserving_order,
)
from src.assembleur_core import TopologyCheminTriplet


class _World:
    def __init__(self):
        self.points = {"C1": (1.0, 0.0), "C2": (0.0, 0.0), "C3": (0.0, 2.0)}
        self.members = {"C1": ["T02:N1", "T01:N0"], "C2": ["T01:N1"], "C3": ["T02:N2"]}
        self.beacons = {"B1": (5.0, 0.0)}

    def getConceptNodeWorldXY(self, node_id, _group_id): return self.points[node_id]
    def getPhysicalNodesForConceptNode(self, node_id): return self.members[node_id]
    def getAtomicNodeLabel(self, node_id): return {"T01:N0": "A", "T01:N1": "O", "T02:N1": "A", "T02:N2": "B"}[node_id]
    def getNodeTypeAtomic(self, node_id): return {"T01:N0": "O", "T01:N1": "B", "T02:N1": "O", "T02:N2": "L"}[node_id]
    def getNodeLabel(self, node_id): return {"C1": "A", "C2": "O", "C3": "B"}[node_id]
    def hasBeacon(self, beacon_id): return beacon_id in self.beacons
    def getBeaconWorldXY(self, beacon_id): return self.beacons[beacon_id]
    def azimutDegFromDxDy(self, dx, dy): return (90.0 - __import__("math").degrees(__import__("math").atan2(dy, dx))) % 360.0


class _Transform:
    calibrated_map = SimpleNamespace(image_size=(1000, 500))
    def world_to_pixel(self, x, y): return (x * 10.0, 100.0 - y * 10.0)
    def world_to_lambert(self, x, y): return (700000.0 + x, 6600000.0 + y)


def _chemins():
    return SimpleNamespace(
        isDefined=True, groupId="G1", orientationUser="cw", pathNodesOrdered=["C1", "C2", "C3"],
        triplets=[TopologyCheminTriplet("C1", "C2", "C3")],
    )


def test_triplets_use_temporary_angles_and_current_world_distances():
    world, chemins = _World(), _chemins()
    source = chemins.triplets[0]
    columns, rows = build_triplets_rows(world, chemins, TripletsExportOptions(include_labels=True, include_angles=True, include_distances=True, reference_beacon_id="B1"))
    assert [name for name, _kind in columns] == ["Triplet", "Az OA (°)", "Az OB (°)", "Angle A-O-B (°)", "Dist OA (km)", "Dist OB (km)"]
    assert rows[0][0] == "A - O - B"
    assert all(isinstance(value, float) for value in rows[0][1:])
    assert source.isGeometrieValide is False
    assert rows[0][-2:] == pytest.approx((1.0, 2.0))


def test_points_follow_path_order_and_aggregate_sorted_physical_members():
    columns, rows = build_points_rows(_World(), _chemins(), PointsExportOptions(include_order=True, include_node_id=True, include_triangles=True, include_cities=True, include_node_types=True, include_pixel=True, pixel_origin_beacon_id="B1", include_lambert=True), map_transform=_Transform())
    assert [name for name, _kind in columns] == ["Ordre", "Node ID", "Triangle(s)", "Ville(s)", "Type(s)", "Pixel X", "Pixel Y", "Lambert X (m)", "Lambert Y (m)"]
    assert len(rows) == 3
    assert rows[0][:5] == [1, "C1", "T01 ; T02", "A", "O"]
    assert rows[0][5:7] == pytest.approx((-40.0, 0.0))
    assert rows[0][7:] == pytest.approx((700001.0, 6600000.0))


def test_points_coordinate_requires_a_map_and_optional_data_is_enforced():
    with pytest.raises(CheminsExportError, match="carte calibrée"):
        build_points_rows(_World(), _chemins(), PointsExportOptions(include_pixel=True, pixel_origin_beacon_id="B1"))
    with pytest.raises(CheminsExportError):
        build_triplets_rows(_World(), _chemins(), TripletsExportOptions())
    assert unique_preserving_order(["A", "B", "A"]) == ["A", "B"]


def test_xlsx_has_one_sheet_numeric_cells_and_optional_references(tmp_path):
    target = export_chemins_xlsx(tmp_path / "points.xlsx", export_kind="points", world=_World(), chemins=_chemins(), scenario_name="Scénario", options=PointsExportOptions(include_order=True, include_pixel=True, pixel_origin_beacon_id="B1", include_reference_info=True), map_transform=_Transform(), map_name="Carte test")
    workbook = openpyxl.load_workbook(target, data_only=False)
    assert workbook.sheetnames == ["Points"]
    sheet = workbook["Points"]
    assert sheet["A1"].value == "Scénario"
    header_row = next(row for row in range(1, sheet.max_row + 1) if sheet.cell(row, 1).value == "Ordre")
    assert sheet.cell(header_row + 1, 2).value == pytest.approx(-40.0)
    assert sheet.cell(header_row + 1, 2).number_format == "0.000"
