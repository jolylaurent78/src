"""Headless Excel export for the current TopologyChemins data."""

from __future__ import annotations

from dataclasses import dataclass
import math
from pathlib import Path
from typing import Callable, Iterable

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from src.assembleur_core import TopologyCheminTriplet
from src.assembleur_catalogue import Catalogue


class CheminsExportError(ValueError):
    """A configuration or data error that can be displayed by the Tk dialog."""


@dataclass(frozen=True)
class TripletsExportOptions:
    include_labels: bool = False
    include_angles: bool = False
    include_distances: bool = False
    reference_beacon_id: str | None = None
    include_reference_info: bool = False

    def __post_init__(self) -> None:
        if self.include_angles and not str(self.reference_beacon_id or "").strip():
            raise CheminsExportError("Une balise de référence est requise pour exporter les angles.")

    @property
    def has_tabular_data(self) -> bool:
        return self.include_labels or self.include_angles or self.include_distances


@dataclass(frozen=True)
class PointsExportOptions:
    include_order: bool = False
    include_node_id: bool = False
    include_triangles: bool = False
    include_cities: bool = False
    include_node_types: bool = False
    include_pixel: bool = False
    pixel_origin_beacon_id: str | None = None
    include_lambert: bool = False
    include_reference_info: bool = False

    def __post_init__(self) -> None:
        if self.include_pixel and not str(self.pixel_origin_beacon_id or "").strip():
            raise CheminsExportError("Une balise d'origine est requise pour exporter les pixels.")

    @property
    def has_tabular_data(self) -> bool:
        return any((self.include_order, self.include_node_id, self.include_triangles,
                    self.include_cities, self.include_node_types, self.include_pixel,
                    self.include_lambert))


@dataclass(frozen=True)
class BeaconsExportOptions:
    include_beacon_id: bool = False
    include_name: bool = False
    include_city_id: bool = False
    include_group: bool = False
    include_order: bool = False
    include_anchor: bool = False
    include_note: bool = False
    include_pixel: bool = False
    pixel_origin_beacon_id: str | None = None
    include_lambert: bool = False
    include_reference_info: bool = False

    def __post_init__(self) -> None:
        if self.include_pixel and not str(self.pixel_origin_beacon_id or "").strip():
            raise CheminsExportError("Une balise d'origine est requise pour exporter les pixels.")

    @property
    def has_tabular_data(self) -> bool:
        return any((self.include_beacon_id, self.include_name, self.include_city_id,
                    self.include_group, self.include_order, self.include_anchor,
                    self.include_note, self.include_pixel, self.include_lambert))


def unique_preserving_order(values: Iterable[str]) -> list[str]:
    seen: set[str] = set()
    result: list[str] = []
    for value in values:
        text = str(value)
        if text not in seen:
            seen.add(text)
            result.append(text)
    return result


def build_triplets_rows(world, chemins, options: TripletsExportOptions) -> tuple[list[tuple[str, str]], list[list[object]]]:
    if not options.has_tabular_data:
        raise CheminsExportError("Sélectionnez au moins une donnée de triplet à exporter.")
    _require_defined(chemins)
    columns: list[tuple[str, str]] = []
    if options.include_labels:
        columns.append(("Triplet", "text"))
    if options.include_angles:
        columns.extend((("Az OA (°)", "angle"), ("Az OB (°)", "angle"), ("Angle A-O-B (°)", "angle")))
    if options.include_distances:
        columns.extend((("Dist OA (km)", "distance"), ("Dist OB (km)", "distance")))

    rows: list[list[object]] = []
    for source in chemins.triplets:
        values: list[object] = []
        if options.include_labels:
            values.append(" - ".join(world.getNodeLabel(node) for node in (source.nodeA, source.nodeO, source.nodeB)))
        if options.include_angles:
            temporary = TopologyCheminTriplet(source.nodeA, source.nodeO, source.nodeB)
            temporary.calculerGeometrie(world, chemins.groupId, chemins.orientationUser, str(options.reference_beacon_id))
            values.extend((_finite(temporary.azOA, "Az OA"), _finite(temporary.azOB, "Az OB"), _finite(temporary.angleDeg, "Angle A-O-B")))
        if options.include_distances:
            point_a = world.getConceptNodeWorldXY(source.nodeA, chemins.groupId)
            point_o = world.getConceptNodeWorldXY(source.nodeO, chemins.groupId)
            point_b = world.getConceptNodeWorldXY(source.nodeB, chemins.groupId)
            values.extend((_distance(point_o, point_a), _distance(point_o, point_b)))
        rows.append(values)
    return columns, rows


def build_points_rows(
    world,
    chemins,
    options: PointsExportOptions,
    *,
    map_transform=None,
    beacon_world_resolver: Callable[[str], tuple[float, float]] | None = None,
) -> tuple[list[tuple[str, str]], list[list[object]]]:
    if not options.has_tabular_data:
        raise CheminsExportError("Sélectionnez au moins une donnée de point à exporter.")
    _require_defined(chemins)
    if (options.include_pixel or options.include_lambert) and map_transform is None:
        raise CheminsExportError("Une carte calibrée est requise pour exporter Pixel ou Lambert-93.")
    columns: list[tuple[str, str]] = []
    if options.include_order:
        columns.append(("Ordre", "integer"))
    if options.include_node_id:
        columns.append(("Node ID", "text"))
    if options.include_triangles:
        columns.append(("Triangle(s)", "text"))
    if options.include_cities:
        columns.append(("Ville(s)", "text"))
    if options.include_node_types:
        columns.append(("Type(s)", "text"))
    if options.include_pixel:
        columns.extend((("Pixel X", "coordinate"), ("Pixel Y", "coordinate")))
    if options.include_lambert:
        columns.extend((("Lambert X (m)", "coordinate"), ("Lambert Y (m)", "coordinate")))

    origin_pixel = None
    if options.include_pixel:
        origin_world = _beacon_world(world, str(options.pixel_origin_beacon_id), beacon_world_resolver)
        origin_pixel = map_transform.world_to_pixel(*origin_world)

    rows: list[list[object]] = []
    for index, node_id in enumerate(chemins.pathNodesOrdered, start=1):
        members = sorted(str(node) for node in world.getPhysicalNodesForConceptNode(node_id))
        point_world = world.getConceptNodeWorldXY(node_id, chemins.groupId)
        values: list[object] = []
        if options.include_order:
            values.append(index)
        if options.include_node_id:
            values.append(str(node_id))
        if options.include_triangles:
            values.append(" ; ".join(unique_preserving_order(_triangle_id(member) for member in members)))
        if options.include_cities:
            values.append(" ; ".join(unique_preserving_order(world.getAtomicNodeLabel(member) for member in members)))
        if options.include_node_types:
            values.append(" ; ".join(unique_preserving_order(world.getNodeTypeAtomic(member) for member in members)))
        if options.include_pixel:
            point_pixel = map_transform.world_to_pixel(*point_world)
            values.extend((_finite(point_pixel[0] - origin_pixel[0], "Pixel X"), _finite(point_pixel[1] - origin_pixel[1], "Pixel Y")))
        if options.include_lambert:
            lambert = map_transform.world_to_lambert(*point_world)
            values.extend((_finite(lambert[0], "Lambert X"), _finite(lambert[1], "Lambert Y")))
        rows.append(values)
    return columns, rows


def build_beacons_rows(catalogue: Catalogue, options: BeaconsExportOptions, *, map_transform=None,
                       beacon_world_resolver: Callable[[str], tuple[float, float]] | None = None) -> tuple[list[tuple[str, str]], list[list[object]]]:
    if not options.has_tabular_data:
        raise CheminsExportError("Sélectionnez au moins une donnée de balise à exporter.")
    if (options.include_pixel or options.include_lambert) and map_transform is None:
        raise CheminsExportError("Une carte calibrée est requise pour exporter Pixel ou Lambert-93.")
    columns: list[tuple[str, str]] = []
    for enabled, label, kind in ((options.include_beacon_id, "ID Balise", "text"), (options.include_name, "Nom", "text"), (options.include_city_id, "ID Ville", "text"), (options.include_group, "Groupe", "text"), (options.include_order, "Ordre", "integer"), (options.include_anchor, "Ancrage", "text"), (options.include_note, "Note", "text")):
        if enabled: columns.append((label, kind))
    if options.include_pixel: columns.extend((("Pixel X", "coordinate"), ("Pixel Y", "coordinate")))
    if options.include_lambert: columns.extend((("Lambert X (m)", "coordinate"), ("Lambert Y (m)", "coordinate")))
    origin_pixel = None
    if options.include_pixel:
        if beacon_world_resolver is None: raise CheminsExportError("Un résolveur de balises est requis pour exporter les pixels.")
        origin_pixel = map_transform.world_to_pixel(*beacon_world_resolver(str(options.pixel_origin_beacon_id)))
    beacons = [beacon for beacon in catalogue.iter_beacons() if not beacon.archived]
    beacons.sort(key=lambda beacon: (not bool(beacon.group), beacon.group.casefold(), beacon.group, beacon.order is None, beacon.order if beacon.order is not None else 0, catalogue.get_city(beacon.city_id).name.casefold(), beacon.beacon_id))
    rows: list[list[object]] = []
    for beacon in beacons:
        city = catalogue.get_city(beacon.city_id); values: list[object] = []
        if options.include_beacon_id: values.append(beacon.beacon_id)
        if options.include_name: values.append(city.name)
        if options.include_city_id: values.append(beacon.city_id)
        if options.include_group: values.append(beacon.group)
        if options.include_order: values.append(beacon.order)
        if options.include_anchor: values.append("Oui" if beacon.usable_as_anchor else "Non")
        if options.include_note: values.append(beacon.note)
        if options.include_pixel or options.include_lambert:
            if beacon_world_resolver is None: raise CheminsExportError("Un résolveur de balises est requis pour exporter les coordonnées.")
            world = beacon_world_resolver(beacon.beacon_id)
        if options.include_pixel:
            pixel = map_transform.world_to_pixel(*world); values.extend((_finite(pixel[0]-origin_pixel[0], "Pixel X"), _finite(pixel[1]-origin_pixel[1], "Pixel Y")))
        if options.include_lambert:
            lambert = map_transform.world_to_lambert(*world); values.extend((_finite(lambert[0], "Lambert X"), _finite(lambert[1], "Lambert Y")))
        rows.append(values)
    return columns, rows


def export_chemins_xlsx(
    file_path: str | Path,
    *,
    export_kind: str,
    world,
    chemins,
    scenario_name: str,
    options: TripletsExportOptions | PointsExportOptions | BeaconsExportOptions,
    catalogue: Catalogue | None = None,
    map_transform=None,
    beacon_world_resolver: Callable[[str], tuple[float, float]] | None = None,
    map_name: str | None = None,
) -> Path:
    target = Path(file_path)
    if target.suffix.lower() != ".xlsx":
        raise CheminsExportError("Le fichier d'export doit avoir l'extension .xlsx.")
    if export_kind == "triplets" and isinstance(options, TripletsExportOptions):
        columns, rows = build_triplets_rows(world, chemins, options)
        references = _triplets_references(scenario_name, chemins, options)
        sheet_name = "Triplets"
    elif export_kind == "points" and isinstance(options, PointsExportOptions):
        columns, rows = build_points_rows(world, chemins, options, map_transform=map_transform, beacon_world_resolver=beacon_world_resolver)
        references = _points_references(scenario_name, chemins, options, map_transform, map_name)
        sheet_name = "Points"
    elif export_kind == "beacons" and isinstance(options, BeaconsExportOptions) and catalogue is not None:
        columns, rows = build_beacons_rows(catalogue, options, map_transform=map_transform, beacon_world_resolver=beacon_world_resolver)
        references = [("Scénario", scenario_name)] + ([("Carte", map_name or "Carte calibrée"), ("Origine pixel", str(options.pixel_origin_beacon_id))] if options.include_pixel else [])
        sheet_name = "Balises"
    else:
        raise CheminsExportError("Type d'export ou options incompatibles.")
    _write_xlsx(target, sheet_name, columns, rows, references if options.include_reference_info else ())
    return target


def _write_xlsx(target: Path, sheet_name: str, columns: list[tuple[str, str]], rows: list[list[object]], references: Iterable[tuple[str, object]]) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = sheet_name
    row_number = 1
    for key, value in references:
        sheet.cell(row=row_number, column=1, value=key)
        sheet.cell(row=row_number, column=2, value=value)
        row_number += 1
    if row_number > 1:
        row_number += 1
    for column_number, (label, _kind) in enumerate(columns, start=1):
        sheet.cell(row=row_number, column=column_number, value=label).font = Font(bold=True)
    for values in rows:
        row_number += 1
        for column_number, value in enumerate(values, start=1):
            cell = sheet.cell(row=row_number, column=column_number, value=value)
            kind = columns[column_number - 1][1]
            if kind == "angle":
                cell.number_format = "0.00"
            elif kind in {"distance", "coordinate"}:
                cell.number_format = "0.000"
    for column_number in range(1, len(columns) + 1):
        width = max(len(str(sheet.cell(row=row, column=column_number).value or "")) for row in range(1, sheet.max_row + 1))
        sheet.column_dimensions[get_column_letter(column_number)].width = min(60, width + 2)
    try:
        workbook.save(target)
    except OSError as exc:
        raise CheminsExportError(f"Impossible d'écrire le fichier Excel : {target}") from exc


def _triplets_references(scenario_name: str, chemins, options: TripletsExportOptions) -> list[tuple[str, object]]:
    values = [("Scénario", scenario_name), ("Orientation du chemin", chemins.orientationUser)]
    if options.include_angles:
        values.append(("Balise de référence", str(options.reference_beacon_id)))
    return values


def _points_references(scenario_name: str, chemins, options: PointsExportOptions, map_transform, map_name: str | None) -> list[tuple[str, object]]:
    values = [("Scénario", scenario_name), ("Orientation du chemin", chemins.orientationUser)]
    if options.include_pixel:
        values.extend((("Carte", map_name or "Carte calibrée"), ("Origine pixel", str(options.pixel_origin_beacon_id)), ("Convention Pixel X", "positif vers la droite"), ("Convention Pixel Y", "positif vers le bas")))
    if options.include_lambert:
        values.append(("Projection", "EPSG:2154"))
    if map_transform is not None and (options.include_pixel or options.include_lambert):
        width, height = map_transform.calibrated_map.image_size
        values.extend((("Largeur image (px)", width), ("Hauteur image (px)", height)))
    return values


def _beacon_world(world, beacon_id: str, resolver: Callable[[str], tuple[float, float]] | None) -> tuple[float, float]:
    point = resolver(beacon_id) if resolver is not None else world.getBeaconWorldXY(beacon_id)
    return (_finite(point[0], "balise.x"), _finite(point[1], "balise.y"))


def _distance(first: tuple[float, float], second: tuple[float, float]) -> float:
    return _finite(math.hypot(float(first[0]) - float(second[0]), float(first[1]) - float(second[1])), "distance")


def _finite(value: object, label: str) -> float:
    try:
        number = float(value)
    except (TypeError, ValueError) as exc:
        raise CheminsExportError(f"{label} doit être numérique.") from exc
    if not math.isfinite(number):
        raise CheminsExportError(f"{label} doit être fini.")
    return number


def _triangle_id(physical_node_id: str) -> str:
    return str(physical_node_id).split(":", 1)[0]


def _require_defined(chemins) -> None:
    if not bool(chemins.isDefined):
        raise CheminsExportError("Aucun chemin défini à exporter.")
